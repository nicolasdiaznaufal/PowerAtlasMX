"""
solar.py - Recurso solar y produccion FV por MWp (Nodal Solar Value).

Que hace
    1. Lee la cache horaria de NASA POWER (data/cache/nasa/celdas/*.json: GHI y T2M por celda de 1 grado).
       -999 / null = dato faltante -> NaN. NUNCA se usa como valor ni se vuelve 0.
    2. Descompone GHI en DNI/DHI (Erbs) y calcula la produccion FV horaria por MWp con pvlib para
       configuracion FIJA y SEGUIDOR de 1 eje N-S con backtracking.
    3. Agrega a baselines: MWh/MWp por mes x hora del dia, anual y factor de planta (FC).

Unidades (todas explicitas en el nombre de la columna/funcion)
    ghi_wm2, dni_wm2, dhi_wm2   W/m2 (promedio de la hora)
    t2m_c                       grados C
    power_w_per_mwp             W por MWp DC instalado (potencia media de la hora, AC)
    energy_wh_per_mwp           Wh por MWp en la hora (= W x 1 h)
    energy_mwh_per_mwp          MWh por MWp en la hora. Numericamente = kWh/kWp = horas-equivalentes.
    cf                          adimensional [0..1] = MWh anuales / (8760 h x 1 MWp)

Supuestos: cada uno marcado [verificado] / [estimado] / [pendiente] en ASSUMPTIONS (abajo).
Capex, opex y tasa NO viven aqui (config/supuestos.yaml, PENDIENTE).

Reloj: pipeline/clock.py (clock-qa). TODO el eje hora/mes/dia de los baselines esta en el reloj comun
CDMX estandar (UTC-6 fijo, sin DST) via `clock.utc_a_cdmx`; la union con el PML usa `clock.ts_cdmx(sistema, fecha, hora)`.
El sistema (SIN/BCS/BCA) es propiedad del nodo y solo importa para esa union; el perfil solar es el de la
ubicacion del nodo expresado en hora CDMX (el mediodia solar de BCA cae ~13-14 h CDMX).
"""
from __future__ import annotations

import json
import math
from dataclasses import dataclass, field, asdict
from pathlib import Path

import numpy as np
import pandas as pd

try:
    import clock
except ImportError:  # importado como pipeline.solar
    from pipeline import clock

MISSING = -999.0
HOURS_PER_YEAR = 8760.0
DAYS_IN_MONTH = {1: 31, 2: 28, 3: 31, 4: 30, 5: 31, 6: 30, 7: 31, 8: 31, 9: 30, 10: 31, 11: 30, 12: 31}

# --------------------------------------------------------------------------------------------
# Reloj: delegado en pipeline/clock.py (clock-qa es el dueno; no editar aqui la convencion).
# --------------------------------------------------------------------------------------------
CLOCK_NAME = "CDMX estandar UTC-6 fijo"             # etiqueta que declara cada archivo de baseline
CDMX_UTC_OFFSET_HOURS = clock.CDMX_UTC_OFFSET_HOURS   # -6, sin DST nunca
STD_UTC_OFFSET_HOURS = clock.STD_UTC_OFFSET_HOURS      # SIN -6, BCS -7, BCA -8 [estimado]; solo para a hora local del sistema
CLOCK_STATUS = clock.CONFIANZA


def utc_to_market_clock(ts_utc: pd.DatetimeIndex, system: str = "SIN") -> pd.DatetimeIndex:
    """UTC inicio de hora -> hora local ESTANDAR del sistema (tz-naive). Solo informativo: los baselines
    usan `utc_to_cdmx`."""
    return clock.utc_a_local_estandar(system, ts_utc)


def utc_to_cdmx(ts_utc: pd.DatetimeIndex) -> pd.DatetimeIndex:
    """UTC inicio de hora -> hora estandar de CDMX (UTC-6 fijo, tz-naive), via clock.utc_a_cdmx."""
    return clock.utc_a_cdmx(ts_utc)


# --------------------------------------------------------------------------------------------
# Supuestos declarados
# --------------------------------------------------------------------------------------------
@dataclass(frozen=True)
class PlantAssumptions:
    # Orientacion
    fixed_azimuth_deg: float = 180.0        # [estimado] hacia el sur (hemisferio norte)
    fixed_tilt_rule: str = "0.76*|lat|+3.1"  # [estimado] regla empirica de inclinacion optima anual
    tracker_axis_azimuth_deg: float = 180.0  # [estimado] eje N-S horizontal
    tracker_max_angle_deg: float = 60.0      # [estimado]
    tracker_gcr: float = 0.35                # [estimado] razon de cobertura para backtracking
    albedo: float = 0.20                     # [estimado] suelo generico
    # Modelos
    decomposition_model: str = "erbs"        # [estimado] GHI -> DNI/DHI
    transposition_model: str = "perez"       # [estimado] allsitescomposite1990
    iam_ashrae_b: float = 0.05               # [estimado] perdida angular sobre la directa
    # Temperatura
    temp_model: str = "sapm_cell/open_rack_glass_glass"   # [estimado]
    wind_speed_ms: float = 1.0               # [estimado] solo si la serie no trae WS2M (cache JSON antigua)
    wind_height_exponent: float = 1.0 / 7.0  # [estimado] ley de potencia 2 m -> 10 m (SAPM espera viento a 10 m)
    use_nasa_dni_dhi: bool = True            # [estimado] DNI/DHI de NASA directos; Erbs si no son validos
    gamma_pdc_per_c: float = -0.0035         # [estimado] coef. de potencia tipico c-Si (igual que legacy)
    # Perdidas (multiplicativas) [estimado: valores tipicos de la industria, sin fuente de proyecto]
    loss_soiling: float = 0.02
    loss_dc_mismatch_wiring: float = 0.03
    loss_availability: float = 0.01
    loss_inverter: float = 0.025             # eficiencia 97.5 % plana
    loss_ac_transformer: float = 0.015
    # Otros
    dc_ac_ratio: float = 1.0                 # [pendiente] sin clipping; ILR real depende del diseno
    degradation_per_year: float = 0.005      # [estimado] 0.5 %/ano; NO aplicada en baseline (ano 1)

    @property
    def system_derate(self) -> float:
        return ((1 - self.loss_soiling) * (1 - self.loss_dc_mismatch_wiring) * (1 - self.loss_availability)
                * (1 - self.loss_inverter) * (1 - self.loss_ac_transformer))


STATUS = {
    "fixed_azimuth_deg": "estimado", "fixed_tilt_rule": "estimado", "tracker_axis": "estimado",
    "tracker_max_angle_deg": "estimado", "tracker_gcr": "estimado", "albedo": "estimado",
    "decomposition_model": "estimado", "transposition_model": "estimado", "iam": "estimado",
    "temp_model": "estimado", "wind": "NASA WS2M real (estimado: extrapolacion 2->10 m, ley 1/7)",
    "dni_dhi": "NASA POWER directo; Erbs de respaldo",
    "gamma_pdc_per_c": "estimado", "losses": "estimado", "dc_ac_ratio": "pendiente",
    "degradation_per_year": "estimado (no aplicada)", "clock": "pendiente: clock-qa",
    "nasa_timestamp": "verificado en legacy (inicio de hora UTC; mejor ajuste con +30 min vs cielo despejado)",
    "nasa_ghi_bias": "pendiente: sin NSRDB disponible para contrastar",
}

DEFAULT = PlantAssumptions()


def degradation_factor(year_index: int, annual_rate: float = DEFAULT.degradation_per_year) -> float:
    """Factor multiplicativo de produccion en el ano `year_index` (1 = primer ano -> 1.0).
    [estimado] degradacion lineal-compuesta; el baseline NO lo aplica."""
    if year_index < 1:
        raise ValueError("year_index >= 1")
    return (1.0 - annual_rate) ** (year_index - 1)


def wh_to_mwh(wh):
    return wh / 1e6


def w_to_wh_hourly(w):
    """Potencia media horaria (W) -> energia en esa hora (Wh), paso de 1 h."""
    return w * 1.0


# --------------------------------------------------------------------------------------------
# 1. Lectura de la cache NASA POWER
# --------------------------------------------------------------------------------------------
def cell_of(lat: float, lon: float) -> tuple[int, int]:
    """Celda de 1 grado (floor lat, floor lon), igual que la cache. [verificado en legacy: GHI identica dentro de la celda]"""
    return math.floor(lat), math.floor(lon)


def cell_path(cache_dir: str | Path, la0: int, lo0: int) -> Path:
    return Path(cache_dir) / "celdas" / f"celda_{la0:+03d}_{lo0:+04d}.json"


def _series(d: dict | None) -> pd.Series:
    if not d:
        return pd.Series(dtype="float64")
    s = pd.Series(d, dtype="float64")          # null -> NaN
    s[s == MISSING] = np.nan                   # -999 -> NaN
    s.index = pd.to_datetime(s.index, format="%Y%m%d%H")
    return s.sort_index()


def load_nasa_cell(path: str | Path) -> pd.DataFrame:
    """DataFrame horario indexado por INICIO de hora en UTC (tz-aware), columnas ghi_wm2, t2m_c.
    Faltantes = NaN. GHI negativa -> NaN (no fisica)."""
    d = json.loads(Path(path).read_text())
    ghi, t = _series(d.get("GHI")), _series(d.get("T2M"))
    df = pd.DataFrame({"ghi_wm2": ghi, "t2m_c": t}).sort_index()
    df.loc[df["ghi_wm2"] < 0, "ghi_wm2"] = np.nan
    df.index = df.index.tz_localize("UTC")
    df.index.name = "ts_utc_hour_start"
    return df


def load_nasa_legacy_point(path: str | Path) -> pd.DataFrame:
    """Archivo por punto (cache/legado: ALLSKY_SFC_SW_DWN, T2M). Mismo formato de salida."""
    d = json.loads(Path(path).read_text())
    df = pd.DataFrame({"ghi_wm2": _series(d.get("ALLSKY_SFC_SW_DWN")), "t2m_c": _series(d.get("T2M"))}).sort_index()
    df.index = df.index.tz_localize("UTC")
    df.index.name = "ts_utc_hour_start"
    return df


CLIMATE_COLS = ("ghi_wm2", "clrsky_ghi_wm2", "dni_wm2", "dhi_wm2", "t2m_c", "ws2m_ms")


def load_climate_cell(path: str | Path) -> pd.DataFrame:
    """Lee una celda de clima. Parquet (cache/nasa/climatologia, ghi/clrsky/dni/dhi/t2m/ws2m) o JSON
    antiguo (solo ghi y t2m). Indice = inicio de hora UTC tz-aware. NaN = faltante (nunca -999)."""
    path = Path(path)
    if path.suffix == ".json":
        return load_nasa_cell(path)
    df = pd.read_parquet(path)
    df = df.replace(MISSING, np.nan)              # defensa: -999 nunca debe ser valor
    for c in CLIMATE_COLS:
        if c in df:
            df.loc[df[c] < 0, c] = np.nan
    if df.index.tz is None:
        df.index = df.index.tz_localize("UTC")
    df.index.name = "ts_utc_hour_start"
    return df.sort_index()


def find_climate_cell(cache_dir: str | Path, la0: int, lo0: int) -> Path:
    """Parquet largo si existe (cache/nasa/climatologia), si no JSON antiguo."""
    name = f"celda_{la0:+03d}_{lo0:+04d}"
    p = Path(cache_dir) / "climatologia" / "celdas" / f"{name}.parquet"
    return p if p.exists() else cell_path(cache_dir, la0, lo0)


# --------------------------------------------------------------------------------------------
# 2. Irradiancia: DNI/DHI
# --------------------------------------------------------------------------------------------
def decompose_ghi(ghi_wm2: pd.Series, zenith_deg: pd.Series, dni_extra_wm2: pd.Series) -> pd.DataFrame:
    """GHI -> (dni, dhi) en W/m2 con Erbs. Horas con sol bajo el horizonte al punto medio pero GHI>0
    (la hora contiene el amanecer/ocaso): toda la GHI se trata como difusa, DNI=0. NaN se propaga."""
    import pvlib
    ok = ghi_wm2.notna()
    g = ghi_wm2.fillna(0.0)
    e = pvlib.irradiance.erbs(g, zenith_deg.clip(upper=90.0), g.index, min_cos_zenith=0.065, max_zenith=87)
    dni, dhi = e["dni"].fillna(0.0), e["dhi"].fillna(0.0)
    below = zenith_deg >= 90.0
    dni = dni.where(~below, 0.0)
    dhi = dhi.where(~below, g)
    dni, dhi = dni.where(ok), dhi.where(ok)
    return pd.DataFrame({"dni_wm2": dni.clip(lower=0), "dhi_wm2": dhi.clip(lower=0)})


def nasa_dni_dhi(ghi_wm2: pd.Series, dni_nasa: pd.Series, dhi_nasa: pd.Series, zenith_deg: pd.Series,
                 dni_extra_wm2: pd.Series) -> pd.DataFrame:
    """DNI/DHI de NASA POWER directos cuando son validos; Erbs en el resto.

    Valido = ambos finitos, >=0, DHI <= 1.05*GHI + 5 W/m2 y GHI finita. Con sol bajo el horizonte al
    punto medio de la hora (zenith>=90, hora con amanecer/ocaso): DNI=0 y DHI=GHI (igual que Erbs).
    Devuelve dni_wm2, dhi_wm2 y `dni_dhi_src` ('nasa' | 'erbs' | 'below_horizon' | 'nodata')."""
    er = decompose_ghi(ghi_wm2, zenith_deg, dni_extra_wm2)
    valid = (ghi_wm2.notna() & dni_nasa.notna() & dhi_nasa.notna() & (dni_nasa >= 0) & (dhi_nasa >= 0)
             & (dhi_nasa <= 1.05 * ghi_wm2 + 5.0))
    below = zenith_deg >= 90.0
    use = valid & ~below
    dni = er["dni_wm2"].where(~use, dni_nasa)
    dhi = er["dhi_wm2"].where(~use, dhi_nasa)
    src = pd.Series("erbs", index=ghi_wm2.index)
    src[use] = "nasa"
    src[below & ghi_wm2.notna()] = "below_horizon"
    src[ghi_wm2.isna()] = "nodata"
    return pd.DataFrame({"dni_wm2": dni, "dhi_wm2": dhi, "dni_dhi_src": src})


# --------------------------------------------------------------------------------------------
# 3. Produccion FV por MWp
# --------------------------------------------------------------------------------------------
def fixed_tilt_deg(lat: float) -> float:
    """[estimado] inclinacion anual optima ~ 0.76*|lat| + 3.1 (grados)."""
    return float(0.76 * abs(lat) + 3.1)


def site_geometry(index_utc: pd.DatetimeIndex, lat: float, lon: float, elevation_m: float = 0.0) -> dict:
    """Posicion solar y extraterrestre al punto medio de cada hora; se calcula una vez por ubicacion."""
    import pvlib
    mid = index_utc + pd.Timedelta(minutes=30)
    sp = pvlib.location.Location(lat, lon, tz="UTC", altitude=elevation_m).get_solarposition(mid, method="nrel_numpy")
    zen, azi, azen = (pd.Series(sp[c].values, index=mid) for c in ("zenith", "azimuth", "apparent_zenith"))
    return {"mid": mid, "zen": zen, "azi": azi, "azen": azen,
            "dni_x": pd.Series(pvlib.irradiance.get_extra_radiation(mid).values, index=mid),
            "am": pd.Series(pvlib.atmosphere.get_relative_airmass(azen).fillna(40.0).values, index=mid)}


def hourly_pv_power(df_utc: pd.DataFrame, lat: float, lon: float, config: str,
                    a: PlantAssumptions = DEFAULT, elevation_m: float = 0.0, geom: dict | None = None,
                    dni_dhi_source: str | None = None) -> pd.DataFrame:
    """Produccion horaria por MWp DC de una ubicacion.

    df_utc: salida de load_climate_cell (ghi_wm2, t2m_c; opcional dni_wm2, dhi_wm2, ws2m_ms).
    config: 'fixed' | 'tracker'. dni_dhi_source: 'nasa' | 'erbs' (por defecto segun a.use_nasa_dni_dhi).
    Devuelve ghi, dni, dhi, poa, t2m, wind, tcell, power_w_per_mwp, energy_wh_per_mwp, energy_mwh_per_mwp,
    dni_dhi_src. Horas sin GHI -> NaN en potencia/energia (no cero). Todo al PUNTO MEDIO de la hora."""
    import pvlib
    if config not in ("fixed", "tracker"):
        raise ValueError("config debe ser 'fixed' o 'tracker'")
    g = geom or site_geometry(df_utc.index, lat, lon, elevation_m)
    mid, zen, azi, azen, dni_x, am = g["mid"], g["zen"], g["azi"], g["azen"], g["dni_x"], g["am"]
    ghi = pd.Series(df_utc["ghi_wm2"].values, index=mid)
    t_air = pd.Series(df_utc["t2m_c"].values, index=mid)
    key = [mid.month, mid.hour]
    if t_air.isna().any():   # T2M faltante: media (mes, hora) de la propia serie
        t_air = t_air.fillna(t_air.groupby(key).transform("mean"))
    if "ws2m_ms" in df_utc and df_utc["ws2m_ms"].notna().any():
        w2 = pd.Series(df_utc["ws2m_ms"].values, index=mid)
        w2 = w2.fillna(w2.groupby(key).transform("mean"))
        wind10 = w2 * (10.0 / 2.0) ** a.wind_height_exponent
    else:
        wind10 = pd.Series(a.wind_speed_ms, index=mid)
    src = dni_dhi_source or ("nasa" if a.use_nasa_dni_dhi else "erbs")
    if src == "nasa" and {"dni_wm2", "dhi_wm2"} <= set(df_utc.columns):
        dec = nasa_dni_dhi(ghi, pd.Series(df_utc["dni_wm2"].values, index=mid),
                           pd.Series(df_utc["dhi_wm2"].values, index=mid), zen, dni_x)
    else:
        dec = decompose_ghi(ghi, zen, dni_x)
        dec["dni_dhi_src"] = np.where(ghi.isna(), "nodata", "erbs")
    dni, dhi = dec["dni_wm2"], dec["dhi_wm2"]

    if config == "fixed":
        tilt = pd.Series(fixed_tilt_deg(lat), index=mid)
        surf_az = pd.Series(a.fixed_azimuth_deg, index=mid)
    else:
        tr = pvlib.tracking.singleaxis(azen, azi, axis_tilt=0, axis_azimuth=a.tracker_axis_azimuth_deg,
                                       max_angle=a.tracker_max_angle_deg, backtrack=True, gcr=a.tracker_gcr)
        tilt = tr["surface_tilt"].fillna(0.0)
        surf_az = tr["surface_azimuth"].fillna(a.tracker_axis_azimuth_deg)

    ghi0, dni0, dhi0 = ghi.fillna(0.0), dni.fillna(0.0), dhi.fillna(0.0)
    poa = pvlib.irradiance.get_total_irradiance(
        tilt, surf_az, zen, azi, dni0, ghi0, dhi0, dni_extra=dni_x, airmass=am,
        albedo=a.albedo, model=a.transposition_model, model_perez="allsitescomposite1990")
    aoi = pvlib.irradiance.aoi(tilt, surf_az, zen, azi)
    iam = pvlib.iam.ashrae(aoi, b=a.iam_ashrae_b).fillna(0.0)
    poa_eff = (poa["poa_direct"].fillna(0).clip(lower=0) * iam
               + poa["poa_diffuse"].fillna(0).clip(lower=0)).clip(lower=0)
    poa_glob = poa["poa_global"].fillna(0).clip(lower=0)
    params = pvlib.temperature.TEMPERATURE_MODEL_PARAMETERS["sapm"]["open_rack_glass_glass"]
    tcell = pvlib.temperature.sapm_cell(poa_glob, t_air, wind10, **params)
    p_w = 1e6 * (poa_eff / 1000.0) * (1.0 + a.gamma_pdc_per_c * (tcell - 25.0)) * a.system_derate
    p_w = p_w.clip(lower=0).where(ghi.notna())
    out = pd.DataFrame({
        "ghi_wm2": ghi.values, "dni_wm2": dni.values, "dhi_wm2": dhi.values, "poa_wm2": poa_glob.values,
        "t2m_c": t_air.values, "wind10_ms": wind10.values, "tcell_c": tcell.values,
        "power_w_per_mwp": p_w.values, "dni_dhi_src": dec["dni_dhi_src"].values,
    }, index=df_utc.index)
    out["energy_wh_per_mwp"] = w_to_wh_hourly(out["power_w_per_mwp"])
    out["energy_mwh_per_mwp"] = wh_to_mwh(out["energy_wh_per_mwp"])
    return out


# --------------------------------------------------------------------------------------------
# 4. Climatologia y baselines
# --------------------------------------------------------------------------------------------
def to_local(df_utc_indexed: pd.DataFrame, system: str = "SIN") -> pd.DataFrame:
    """Reindexa a hora local estandar del sistema (clock.utc_a_local_estandar). Agrega date, month, hour, year."""
    out = df_utc_indexed.copy()
    out.index = utc_to_market_clock(out.index, system)
    out.index.name = "ts_local_hour_start"
    out["date"] = out.index.normalize()
    out["month"] = out.index.month
    out["hour"] = out.index.hour
    out["year"] = out.index.year
    return out


def to_cdmx(df_utc_indexed: pd.DataFrame) -> pd.DataFrame:
    """Reindexa a hora estandar de CDMX (clock.utc_a_cdmx; igual para SIN/BCS/BCA). Agrega date, month, hour,
    year tomados del reloj CDMX. Es el eje de todos los baselines."""
    out = df_utc_indexed.copy()
    out.index = utc_to_cdmx(out.index)
    out.index.name = "ts_cdmx_hour_start"
    out["date"] = out.index.normalize()
    out["month"] = out.index.month
    out["hour"] = out.index.hour
    out["year"] = out.index.year
    return out


def restrict_years(local: pd.DataFrame, years) -> pd.DataFrame:
    """Solo anos locales completos pedidos (evita sesgo estacional de un ano parcial)."""
    return local if years is None else local[local["year"].isin(list(years))]


def valid_day_mask(local: pd.DataFrame, col: str = "ghi_wm2") -> pd.Series:
    """True en cada hora cuyo DIA (del reloj del indice: CDMX en los baselines) tiene las 24 horas con dato. Dias incompletos se
    excluyen (no se cuentan como dias sin sol)."""
    ok = local[col].notna()
    return ok.groupby(local["date"]).transform("sum") == 24


def climatology_month_hour(local: pd.DataFrame, cols=("ghi_wm2", "dni_wm2", "dhi_wm2", "t2m_c", "wind10_ms"),
                           years=None) -> pd.DataFrame:
    """Media por (mes, hora local) sobre dias validos de los anos pedidos; sufijo _mean + n_days."""
    v = restrict_years(local, years)
    v = v[valid_day_mask(v)]
    cols = [c for c in cols if c in v]
    g = v.groupby(["month", "hour"])
    res = g[cols].mean().add_suffix("_mean")
    res["n_days"] = g["date"].nunique()
    return res.reset_index()


def baseline_month_hour(local: pd.DataFrame, energy_col: str = "energy_mwh_per_mwp", years=None) -> pd.DataFrame:
    """MWh/MWp por (mes, hora) del DIA PROMEDIO sobre dias validos, mas n_days."""
    v = restrict_years(local, years)
    v = v[valid_day_mask(v)]
    g = v.groupby(["month", "hour"])
    res = g[energy_col].mean().rename("mwh_per_mwp_mean_day").to_frame()
    res["n_days"] = g["date"].nunique()
    return res.reset_index()


def annual_by_year(local: pd.DataFrame, energy_col: str = "energy_mwh_per_mwp", min_valid_days: int = 350) -> pd.DataFrame:
    """MWh/MWp-ano por ano local: suma de dias validos x 365/n_dias; NaN si hay < min_valid_days (ano parcial).
    Permite peor ano / mediana / mejor ano."""
    v = local[valid_day_mask(local)]
    g = v.groupby("year")
    nd = g["date"].nunique()
    tot = g[energy_col].sum() * 365.0 / nd
    res = pd.DataFrame({"mwh_per_mwp_year": tot, "valid_days": nd})
    res.loc[res["valid_days"] < min_valid_days, "mwh_per_mwp_year"] = np.nan
    return res.reset_index()


def annual_from_month_hour(mh: pd.DataFrame, min_days_per_month: int = 5) -> dict:
    """Anual (MWh/MWp-ano) = sum_mes DIAS_MES x sum_hora(media dia). Si algun mes tiene < min_days
    dias validos, el anual es NaN (no se inventa). Ano de 365 d. cf = anual / 8760."""
    nd = mh.groupby("month")["n_days"].max()
    monthly = mh.groupby("month")["mwh_per_mwp_mean_day"].sum()
    missing = [m for m in range(1, 13) if m not in nd.index or nd[m] < min_days_per_month]
    monthly_mwh = {m: float(monthly[m] * DAYS_IN_MONTH[m]) for m in monthly.index}
    annual = float("nan") if missing else float(sum(monthly_mwh.values()))
    return {"annual_mwh_per_mwp": annual, "cf": annual / HOURS_PER_YEAR if not missing else float("nan"),
            "months_missing": missing, "monthly_mwh_per_mwp": monthly_mwh,
            "valid_days": int(mh.groupby("month")["n_days"].max().sum())}


def window_yield(local: pd.DataFrame, start: str, end: str, energy_col: str = "energy_mwh_per_mwp") -> dict:
    """Rendimiento anualizado en una ventana [start, end] (fechas locales) con dias validos,
    escalado x365/n_dias (misma convencion que el legacy valor_nodo.py). MWh/MWp-ano (= kWh/kWp-ano)."""
    m = valid_day_mask(local) & (local["date"] >= pd.Timestamp(start)) & (local["date"] <= pd.Timestamp(end))
    n = local.loc[m, "date"].nunique()
    if n == 0:
        return {"yield_mwh_per_mwp_year": float("nan"), "valid_days": 0}
    return {"yield_mwh_per_mwp_year": float(local.loc[m, energy_col].sum() * 365.0 / n), "valid_days": int(n)}


def sine_shape(hours=range(24), sunrise: float = 6.0, sunset: float = 18.0) -> np.ndarray:
    """Perfil 'seno' normalizado (suma=1) sobre el reloj CDMX estandar: semiciclo de seno entre `sunrise`=6 y
    `sunset`=18 h CDMX (centro 12:00 CDMX), en el punto medio de cada hora. Es el MISMO para todos los nodos:
    no se ajusta a la longitud (en SIN coincide con la definicion anterior en hora local; para BCS/BCA el
    centro 12:00 CDMX = 11:00/10:00 local). [estimado] definicion propia. No mostrar sin aviso."""
    h = np.asarray(list(hours), dtype=float) + 0.5
    y = np.where((h > sunrise) & (h < sunset), np.sin(np.pi * (h - sunrise) / (sunset - sunrise)), 0.0)
    return y / y.sum()


def align_production_pml(prod_cdmx: pd.DataFrame, pml: pd.DataFrame, system: str,
                         energy_col: str = "energy_mwh_per_mwp") -> pd.DataFrame:
    """Une produccion (indice = inicio de hora en CDMX estandar, salida de `to_cdmx`) con el PML publicado
    (fecha, hora 1..n) por instante usando clock.ts_cdmx(sistema, fecha, hora). Los dias de 23/25 horas se
    conservan (cada hora publicada tiene su propio instante). Devuelve ts_cdmx, energy, pml, fecha, hora."""
    p = pml.copy()
    p["ts_cdmx"] = clock.ts_cdmx(system, p["fecha"], p["hora"]).values
    e = prod_cdmx[[energy_col]].rename_axis("ts_cdmx").reset_index()
    return e.merge(p[["ts_cdmx", "fecha", "hora", "pml"]], on="ts_cdmx", how="inner")


# --------------------------------------------------------------------------------------------
# 5. Corrida por ubicacion + metadatos
# --------------------------------------------------------------------------------------------
BASELINE_YEARS = tuple(range(2015, 2026))   # anos completos con GHI (COVERAGE.json); 2026 es parcial y se excluye


def nasa_cell_metadata(df_utc: pd.DataFrame, cell: tuple[int, int] | None = None) -> dict:
    ok = df_utc["ghi_wm2"].notna()
    return {
        "source": "NASA POWER hourly (cache local)",
        "parameters_available": [c for c in CLIMATE_COLS if c in df_utc],
        "time_standard": "UTC, inicio de hora",
        "period_available": [str(df_utc.index.min()), str(df_utc.index.max())],
        "period_with_ghi": [str(df_utc.index[ok].min()), str(df_utc.index[ok].max())] if ok.any() else None,
        "hours_total": int(len(df_utc)), "hours_with_ghi": int(ok.sum()),
        "pct_missing": float(100 * (1 - ok.mean())),
        "cell": list(cell) if cell else None,
    }


def run_location(df_utc: pd.DataFrame, lat: float, lon: float, system: str = "SIN",
                 a: PlantAssumptions = DEFAULT, years=BASELINE_YEARS, compare_erbs: bool = False) -> dict:
    """Para una ubicacion: climatologia mes-hora y, por config (fixed/tracker), baseline mes-hora,
    anual (sobre `years`), serie anual por ano; todo en hora CDMX estandar (`system` queda por compatibilidad). compare_erbs: tracker con Erbs para medir la diferencia."""
    geom = site_geometry(df_utc.index, lat, lon)
    out = {}
    for cfg in ("fixed", "tracker"):
        pv = hourly_pv_power(df_utc, lat, lon, cfg, a, geom=geom)
        loc = to_cdmx(pv)
        mh = baseline_month_hour(loc, years=years)
        out[cfg] = {"hourly_local": loc, "month_hour": mh, "by_year": annual_by_year(loc),
                    **annual_from_month_hour(mh)}
    out["climatology"] = climatology_month_hour(out["tracker"]["hourly_local"], years=years)
    if compare_erbs:
        pv = hourly_pv_power(df_utc, lat, lon, "tracker", a, geom=geom, dni_dhi_source="erbs")
        mh = baseline_month_hour(to_cdmx(pv), years=years)
        out["tracker_erbs"] = annual_from_month_hour(mh)
    return out


# --------------------------------------------------------------------------------------------
# 6. Corrida por lote (CLI): nodos y celdas -> data/baselines/
# --------------------------------------------------------------------------------------------
LEGACY_WINDOWS = {"Y1": ("2024-08-01", "2025-07-31"), "Y2": ("2025-08-01", "2026-07-31")}


def load_node_coords(root: str | Path) -> pd.DataFrame:
    """Coordenadas por nodo. Devuelve nodo, sistema, lat, lon, coord_source, coord_precision.

    1. 'municipio_catalogo': centroide del MUNICIPIO del catalogo CENACE (INEGI), no la subestacion; mismo
       metodo que pipeline/legacy/geolocalizar_nodos.py. SIN, BCA y BCS.
    2. 'renombre:<codigo_nuevo>': el nodo ya no esta en el catalogo pero la hoja 'Control de Cambios' del
       mismo catalogo declara "Antes X ahora Y"; se usa la ubicacion (municipio) de Y. Misma subestacion.
    No hay otras fuentes: bajas sin renombre, interconexiones con el extranjero ('No Aplica') y municipios sin
    centroide en municipios_centroides.csv quedan SIN coordenadas (no se inventan).
    Si falta el catalogo/openpyxl cae a legacy/outputs/deck_coords.json (coord_source='deck_coords')."""
    root = Path(root)
    try:
        import openpyxl
        xl = next((root / "data" / "ref").glob("Cat*NodosP*.xlsx"))
        wb = openpyxl.load_workbook(xl, read_only=True, data_only=True)
        rows = [r for r in wb.worksheets[0].iter_rows(min_row=3, values_only=True) if r and r[3]]
        cat = pd.DataFrame([(r[0], r[3], r[14], r[16]) for r in rows], columns=["sistema", "nodo", "ent", "mun"])
        z = lambda s_, n: pd.to_numeric(s_, errors="coerce").apply(lambda x: str(int(x)).zfill(n) if pd.notna(x) else None)
        cat["CVE_ENT"], cat["CVE_MUN"] = z(cat["ent"], 2), z(cat["mun"], 3)
        cen = pd.read_csv(root / "data" / "ref" / "municipios_centroides.csv", dtype={"CVE_ENT": str, "CVE_MUN": str})
        j = cat.merge(cen[["CVE_ENT", "CVE_MUN", "LAT_DEC", "LON_DEC"]], on=["CVE_ENT", "CVE_MUN"], how="left")
        j = j.dropna(subset=["LAT_DEC"]).rename(columns={"LAT_DEC": "lat", "LON_DEC": "lon"})
        j["coord_source"], j["coord_precision"] = "municipio_catalogo", "centroide municipal"
        out = j[["nodo", "sistema", "lat", "lon", "coord_source", "coord_precision"]].drop_duplicates("nodo")
        # renombres declarados en 'Control de Cambios'
        import re
        ren = {}
        if len(wb.worksheets) > 1:
            for r in wb.worksheets[1].iter_rows(values_only=True):
                for c in r:
                    if isinstance(c, str):
                        for m in re.finditer(r"antes[:\s]+(\d\d[A-Z0-9]{3}-\d+)\s+ahora[:\s]+(\d\d[A-Z0-9]{3}-\d+)", c, re.I):
                            ren[m.group(1).upper()] = m.group(2).upper()
        have = out.set_index("nodo")
        add = []
        for old, new in ren.items():
            if old not in have.index and new in have.index:
                n = have.loc[new]
                add.append({"nodo": old, "sistema": n["sistema"], "lat": n["lat"], "lon": n["lon"],
                            "coord_source": f"renombre:{new}", "coord_precision": "centroide municipal (del nodo renombrado)"})
        if add:
            out = pd.concat([out, pd.DataFrame(add)], ignore_index=True)
        return out.reset_index(drop=True)
    except (ImportError, StopIteration, FileNotFoundError):
        c = json.loads((root / "legacy" / "outputs" / "deck_coords.json").read_text())
        return pd.DataFrame([(n, "SIN", v[0], v[1], "deck_coords", "centroide municipal") for n, v in c.items()],
                            columns=["nodo", "sistema", "lat", "lon", "coord_source", "coord_precision"])


def _worker(args):
    loc_id, loc_type, system, lat, lon, cache_dir, coord_source = args
    la0, lo0 = cell_of(lat, lon)
    df = load_climate_cell(find_climate_cell(cache_dir, la0, lo0))
    r = run_location(df, lat, lon, system, compare_erbs=True)
    ann, mh, wy, yr = [], [], [], []
    for cfg in ("fixed", "tracker"):
        c = r[cfg]
        loc = c["hourly_local"]
        by = c["by_year"]
        yy = by[by["year"].isin(BASELINE_YEARS)]["mwh_per_mwp_year"].dropna()
        row = {"loc_id": loc_id, "loc_type": loc_type, "system": system, "lat": lat, "lon": lon, "cell_lat": la0,
               "cell_lon": lo0, "coord_source": coord_source, "config": cfg, "annual_mwh_per_mwp": c["annual_mwh_per_mwp"], "cf": c["cf"],
               "valid_days": c["valid_days"], "months_missing": ",".join(map(str, c["months_missing"])),
               "years_used": len(yy), "annual_year_min": yy.min() if len(yy) else np.nan,
               "annual_year_p10": yy.quantile(0.10) if len(yy) else np.nan,
               "annual_year_median": yy.median() if len(yy) else np.nan,
               "annual_year_max": yy.max() if len(yy) else np.nan}
        if cfg == "tracker":
            e = r["tracker_erbs"]["annual_mwh_per_mwp"]
            row["annual_erbs_mwh_per_mwp"] = e
            src = loc.loc[loc["year"].isin(BASELINE_YEARS), "dni_dhi_src"]
            row["pct_hours_nasa_dni_dhi"] = float((src == "nasa").sum() / max((src.isin(["nasa", "erbs"])).sum(), 1) * 100)
        ann.append(row)
        m = c["month_hour"].copy()
        m.insert(0, "config", cfg)
        m.insert(0, "loc_id", loc_id)
        mh.append(m)
        b = by.copy()
        b.insert(0, "config", cfg)
        b.insert(0, "loc_id", loc_id)
        yr.append(b)
        for w, (s, e_) in LEGACY_WINDOWS.items():
            wy.append({"loc_id": loc_id, "config": cfg, "window": w, "assumptions": "default",
                       **window_yield(loc, s, e_)})
    # tracker sin perdidas de sistema (comparacion con legacy): potencia = potencia / derate (multiplicativo)
    loc0 = r["tracker"]["hourly_local"].copy()
    loc0["energy_mwh_per_mwp"] = loc0["energy_mwh_per_mwp"] / DEFAULT.system_derate
    for w, (s, e_) in LEGACY_WINDOWS.items():
        wy.append({"loc_id": loc_id, "config": "tracker", "window": w, "assumptions": "sin_perdidas_sistema",
                   **window_yield(loc0, s, e_)})
    cl = r["climatology"].copy()
    cl.insert(0, "loc_id", loc_id)
    return pd.DataFrame(ann), pd.concat(mh), pd.DataFrame(wy), cl, pd.concat(yr)


def run_batch(root: str | Path, out_dir: str | Path | None = None, n_workers: int | None = None,
              limit: int | None = None) -> dict:
    from concurrent.futures import ProcessPoolExecutor
    root = Path(root)
    cache_dir = root / "data" / "cache" / "nasa"
    out_dir = Path(out_dir or root / "data" / "baselines")
    out_dir.mkdir(parents=True, exist_ok=True)
    nodes = load_node_coords(root)
    nodes["cell_ok"] = [find_climate_cell(cache_dir, *cell_of(la, lo)).exists() for la, lo in zip(nodes.lat, nodes.lon)]
    skipped = nodes[~nodes["cell_ok"]]
    nodes = nodes[nodes["cell_ok"]]
    jobs = [(r.nodo, "node", r.sistema, float(r.lat), float(r.lon), str(cache_dir), r.coord_source) for r in nodes.itertuples()]
    cells = []
    for f in sorted((cache_dir / "climatologia" / "celdas").glob("celda_*.parquet")):
        la0, lo0 = int(f.stem[6:9]), int(f.stem[10:])
        cells.append((f"cell_{la0:+03d}_{lo0:+04d}", "cell", "SIN", la0 + 0.5, lo0 + 0.5, str(cache_dir), "centro_celda"))
    jobs = (cells + jobs)[:limit] if limit else cells + jobs
    with ProcessPoolExecutor(n_workers) as ex:
        res = list(ex.map(_worker, jobs, chunksize=8))
    ann = pd.concat([r[0] for r in res], ignore_index=True)
    mh = pd.concat([r[1] for r in res], ignore_index=True)
    wy = pd.concat([r[2] for r in res], ignore_index=True)
    cl = pd.concat([r[3] for r in res], ignore_index=True)
    yr = pd.concat([r[4] for r in res], ignore_index=True)
    for _df in (ann, mh, wy, cl, yr):
        _df["clock"] = CLOCK_NAME
    ann.to_parquet(out_dir / "solar_baselines.parquet", index=False)
    mh.to_parquet(out_dir / "solar_baselines_month_hour.parquet", index=False)
    wy.to_parquet(out_dir / "solar_window_yields.parquet", index=False)
    cl.to_parquet(out_dir / "solar_climatology_month_hour.parquet", index=False)
    yr.to_parquet(out_dir / "solar_annual_by_year.parquet", index=False)
    cov = {}
    cp = cache_dir / "climatologia" / "COVERAGE.json"
    if cp.exists():
        c = json.loads(cp.read_text())
        cov = {k: c.get(k) for k in ("generated", "source", "variables", "n_cells")}
    sample = load_climate_cell(find_climate_cell(cache_dir, 29, -111))
    meta = {
        "generated_by": "pipeline/solar.py",
        "units": {"annual_mwh_per_mwp": "MWh/MWp-ano (AC, por MWp DC)", "cf": "adimensional = MWh/(8760 h x MWp)",
                  "mwh_per_mwp_mean_day": "MWh/MWp por hora del dia promedio del mes (hora CDMX estandar UTC-6 fijo, inicio de hora)",
                  "annual_year_*": "MWh/MWp-ano por ano calendario local (min, p10, mediana, max sobre anos completos)"},
        "climate": {**nasa_cell_metadata(sample, (29, -111)), "coverage": cov,
                    "baseline_years": [BASELINE_YEARS[0], BASELINE_YEARS[-1]], "n_baseline_years": len(BASELINE_YEARS),
                    "partial_2026": "excluido del baseline (Ene-Jun: sesgo estacional); solo en ventanas legacy"},
        "assumptions": asdict(DEFAULT), "assumption_status": STATUS, "system_derate": DEFAULT.system_derate,
        "clock": CLOCK_NAME,
        "clock_detail": {"module": "pipeline/clock.py", "profile_axis": "month/hour/date/year = clock.utc_a_cdmx(ts_utc)",
                         "pml_join": "clock.ts_cdmx(sistema, fecha, hora)", "cdmx_utc_offset_hours": CDMX_UTC_OFFSET_HOURS,
                         "desfase_a_cdmx_horas": {s_: clock.desfase_a_cdmx_horas(s_) for s_ in clock.SISTEMAS},
                         "profile_is_per_node_location": True, "system_is_node_property": True,
                         "confidence": dict(clock.CONFIANZA)},
        "n_locations": {"node": int((ann["loc_type"] == "node").sum() // 2),
                        "cell": int((ann["loc_type"] == "cell").sum() // 2),
                        "node_by_system": ann[ann["loc_type"] == "node"].groupby("system")["loc_id"].nunique().to_dict()},
        "nodes_skipped_no_climate_cell": {
            "count": int(len(skipped)), "by_system": skipped["sistema"].value_counts().to_dict(),
            "cells_needed": sorted([list(c) for c in {cell_of(la, lo) for la, lo in zip(skipped.lat, skipped.lon)}])},
        "coords_note": "lat/lon = centroide del municipio (catalogo CENACE + INEGI), no la subestacion; ver coord_source",
        "coord_sources": ann[ann["loc_type"] == "node"].groupby("coord_source")["loc_id"].nunique().to_dict(),
        "no_nsrdb": True,
    }
    (out_dir / "solar_baselines_metadata.json").write_text(json.dumps(meta, indent=2, default=str))
    return meta


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", default=str(Path(__file__).resolve().parents[1]))
    ap.add_argument("--limit", type=int, default=None)
    ap.add_argument("--workers", type=int, default=None)
    a = ap.parse_args()
    m = run_batch(a.root, limit=a.limit, n_workers=a.workers)
    print(json.dumps(m["n_locations"]))
